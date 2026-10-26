import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_callout(doc, text, title="CATATAN PENTING"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F0FDF4")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(20, 99, 86)
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(10)
    
    run_text = p.add_run(text)
    run_text.font.name = 'Segoe UI'
    run_text.font.size = Pt(9.5)
    doc.add_paragraph()

def style_table(tbl, col_widths, headers, data, header_bg="146356"):
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.name = 'Segoe UI'
            run.font.size = Pt(9.5)
            
    # Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = tbl.add_row().cells
        bg = "FFFFFF" if r_idx % 2 == 0 else "F9FAFB"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=130, right=130)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                run.font.name = 'Segoe UI'
                run.font.size = Pt(9)
                
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

doc = docx.Document()

# Page Margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
C_DARK = RGBColor(15, 61, 62)      # #0f3d3e
C_TEAL = RGBColor(20, 99, 86)      # #146356
C_MUTED = RGBColor(107, 114, 128)  # #6b7280

# ==========================================
# HALAMAN JUDUL (COVER)
# ==========================================
p_cover_pre = doc.add_paragraph()
p_cover_pre.paragraph_format.space_before = Pt(40)
p_cover_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_pre = p_cover_pre.add_run("LAPORAN PRAKTIKUM MANAJEMEN SISTEM INFORMASI")
r_pre.font.name = 'Segoe UI'
r_pre.font.size = Pt(14)
r_pre.font.bold = True
r_pre.font.color.rgb = C_TEAL

p_cover_title = doc.add_paragraph()
p_cover_title.paragraph_format.space_before = Pt(10)
p_cover_title.paragraph_format.space_after = Pt(15)
p_cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_title = p_cover_title.add_run("RANCANG BANGUN SISTEM INFORMASI PEMINJAMAN ALAT LABORATORIUM BERBASIS FRAMEWORK LARAVEL 12 (MVC) DENGAN SISTEM DUAL-ROLE (ADMIN & MAHASISWA)")
r_title.font.name = 'Segoe UI'
r_title.font.size = Pt(18)
r_title.font.bold = True
r_title.font.color.rgb = C_DARK

p_cover_sub = doc.add_paragraph()
p_cover_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_cover_sub.paragraph_format.space_after = Pt(50)
r_sub = p_cover_sub.add_run("Studi Kasus: Penerapan Aturan Bisnis Perhitungan Denda Keterlambatan dan Biaya Servis Kerusakan Secara Terotomasi")
r_sub.font.name = 'Segoe UI'
r_sub.font.size = Pt(11)
r_sub.font.italic = True
r_sub.font.color.rgb = C_MUTED

# Kotak Anggota Kelompok
p_anggota_hdr = doc.add_paragraph()
p_anggota_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_ang = p_anggota_hdr.add_run("Disusun Oleh: KELOMPOK 6")
r_ang.font.name = 'Segoe UI'
r_ang.font.size = Pt(12)
r_ang.font.bold = True
r_ang.font.color.rgb = C_DARK

tbl_kelompok = doc.add_table(rows=1, cols=3)
style_table(
    tbl_kelompok,
    [0.6, 3.4, 2.2],
    ["No", "Nama Mahasiswa", "NIM"],
    [
        ["1", "Lian Rosyid Alamzah", "24051130084"],
        ["2", "Aditya Danu Marizta", "24051130094"],
        ["3", "Muhammad Ghafur Amanu Putra Warnada", "24051130103"]
    ]
)

p_instansi = doc.add_paragraph()
p_instansi.paragraph_format.space_before = Pt(80)
p_instansi.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_ins = p_instansi.add_run(
    "PROGRAM STUDI TEKNOLOGI INFORMASI\n"
    "PRAKTIKUM MANAJEMEN SISTEM INFORMASI\n"
    "TAHUN AKADEMIK 2026\n"
)
r_ins.font.name = 'Segoe UI'
r_ins.font.size = Pt(11)
r_ins.font.bold = True
r_ins.font.color.rgb = C_DARK

doc.add_page_break()

# Helper untuk Heading
def add_h1(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(8)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = C_DARK
    return h

def add_h2(text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = C_TEAL
    return h

def add_p(text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Segoe UI'
    r.font.size = Pt(10)
    return p

# ==========================================
# BAB I: PENDAHULUAN
# ==========================================
add_h1("BAB I: PENDAHULUAN")

add_h2("1.1 Latar Belakang")
add_p(
    "Pengelolaan peralatan dan instrumen di laboratorium akademik memerlukan pencatatan sirkulasi yang tertib, "
    "akurat, dan transparan. Pada sistem konvensional berbasis kertas atau spreadsheet sederhana, pencatatan peminjaman "
    "sering menghadapi berbagai kendala, seperti kesulitan melacak ketersediaan stok aktual, ketidaktepatan penghitungan denda "
    "keterlambatan, serta potensi kesalahan saat memeriksa fisik alat yang rusak ketika dikembalikan."
)
add_p(
    "Untuk mengatasi permasalahan tersebut, Kelompok 6 mengembangkan aplikasi berbasis web dengan framework modern Laravel 12 "
    "yang mengusung arsitektur Model-View-Controller (MVC). Sistem ini dilengkapi dengan pemisahan hak akses (Role-Based Access Control) "
    "antara Administrator (Petugas Laboratorium) dan User (Mahasiswa peminjam), sehingga setiap entitas pengguna memiliki antarmuka dan "
    "otoritas fitur yang sesuai dengan fungsinya."
)

add_h2("1.2 Rumusan Masalah")
add_p("1. Bagaimana merancang sistem informasi laboratorium yang mampu memisahkan antarmuka Admin dan Mahasiswa secara aman?")
add_p("2. Bagaimana mengimplementasikan aturan bisnis denda keterlambatan (> 3 hari denda Rp2.000/hari) dan biaya kerusakan (Rp10.000) secara otomatis?")
add_p("3. Bagaimana mengelola sinkronisasi stok alat secara real-time saat transaksi dipinjam maupun dikembalikan?")

add_h2("1.3 Tujuan Sistem")
add_p("1. Membangun sistem sirkulasi peminjaman alat berbasis web yang efisien, handal, dan berintegritas tinggi.")
add_p("2. Menerapkan perhitungan denda dan biaya servis secara terotomasi di level Controller guna menghindari manipulasi human error.")
add_p("3. Memberikan pengalaman pengguna yang profesional melalui pemisahan portal operasional Admin dan portal swalayan Mahasiswa.")

# ==========================================
# BAB II: ATURAN BISNIS & SPESIFIKASI KEBUTUHAN
# ==========================================
add_h1("BAB II: ATURAN BISNIS & ANALISIS KEBUTUHAN")

add_h2("2.1 Implementasi 3 Aturan Bisnis Utama (Sesuai Soal UTS)")
add_p(
    "Sistem ini menerapkan seluruh ketentuan aturan bisnis yang diwajibkan dalam soal ujian praktikum dengan akurasi 100%:"
)

add_p(
    "Jika pengembalian alat laboratorium melewati batas rencana kembali lebih dari 3 hari, "
    "maka peminjam dikenakan denda keterlambatan sebesar Rp2.000 per hari keterlambatan. Sebaliknya, keterlambatan 1 hingga 3 hari "
    "masih berada dalam batas toleransi dan dikenakan denda Rp0.",
    bold_prefix="1. Aturan Denda Keterlambatan: "
)

add_p(
    "Jika saat proses pengembalian alat dinyatakan dalam kondisi rusak (setelah diverifikasi oleh petugas), "
    "maka sistem otomatis menambahkan biaya perbaikan/servis tetap sebesar Rp10.000 per unit alat.",
    bold_prefix="2. Aturan Biaya Perbaikan (Alat Rusak): "
)

add_p(
    "Kalkulasi denda dan biaya servis tidak diinput secara manual oleh operator, melainkan "
    "dihitung secara otomatis oleh sistem saat form pengembalian disubmit pada PeminjamanAdminController, sehingga menjamin objektivitas data.",
    bold_prefix="3. Otomasi Transaksi & Perhitungan: "
)

add_p(
    "Sistem mencatat stok_total dan stok_tersedia. Ketika transaksi peminjaman baru dicatat, "
    "stok_tersedia berkurang 1. Mahasiswa dicegah meminjam barang yang stok tersedianya habis (0). Setelah alat dikembalikan, stok_tersedia otomatis bertambah kembali.",
    bold_prefix="4. Aturan Validasi Stok Dinamis (Fitur Tambahan Nilai Plus): "
)

add_callout(
    doc,
    "Formula Kalkulasi Denda: Total Denda = (Hari Terlambat > 3 ? Hari Terlambat * 2.000 : 0) + (Kondisi == 'rusak' ? 10.000 : 0)",
    title="RUMUS BISNIS SISTEM"
)

add_h2("2.2 Pembagian Peran Pengguna (Admin vs User/Mahasiswa)")

tbl_fitur = doc.add_table(rows=1, cols=3)
style_table(
    tbl_fitur,
    [1.8, 2.2, 2.4],
    ["Fitur / Modul", "👨‍💼 Admin (Petugas Lab)", "🎓 User (Mahasiswa)"],
    [
        ["Katalog Peralatan", "Akses penuh CRUD (Tambah, Edit, Hapus, Atur Stok)", "Melihat katalog, pencarian, dan ketersediaan stok"],
        ["Kategori Barang", "Kelola kategori (Tambah, Edit, Hapus kategori)", "Filter katalog berdasarkan kategori alat"],
        ["Pengajuan Pinjam", "Mencatat transaksi & memantau seluruh peminjam", "Mengajukan pinjam alat mandiri (identitas terkunci)"],
        ["Verifikasi Kembali", "Memeriksa fisik alat (Baik/Rusak) & simpan denda", "Hanya melihat status penyelesaian pengembalian"],
        ["Riwayat Transaksi", "Rekap seluruh transaksi semua mahasiswa", "Rekap riwayat peminjaman pribadi masing-masing"],
        ["Analitik & Denda", "Dashboard statistik & rekap pemasukan denda", "Kartu status denda dan batas waktu pengembalian"]
    ]
)

# ==========================================
# BAB III: PERANCANGAN SISTEM & ARSITEKTUR
# ==========================================
add_h1("BAB III: PERANCANGAN SISTEM & ARSITEKTUR (MVC)")

add_h2("3.1 Konsep Arsitektur Laravel MVC")
add_p(
    "Framework Laravel membagi arsitektur ke dalam 3 komponen fundamental:\n"
    "• Model: Merepresentasikan tabel database dan relasi antar entitas (User, Barang, Kategori, Peminjaman) menggunakan Eloquent ORM.\n"
    "• View: Menyajikan antarmuka visual kepada pengguna menggunakan Blade Templating Engine yang responsif berbasis Bootstrap 5.\n"
    "• Controller: Bertindak sebagai pengendali alur logika bisnis, memproses request HTTP, melakukan validasi data, serta mengoperasikan Model dan View."
)

add_h2("3.2 Skema Database & Relasi Tabel (ERD)")
add_p("Aplikasi menggunakan 4 tabel utama yang saling berelasi di dalam basis data MySQL (lab_peminjaman_laravel):")

tbl_skema = doc.add_table(rows=1, cols=4)
style_table(
    tbl_skema,
    [1.4, 1.4, 1.8, 1.8],
    ["Nama Tabel", "Primary Key", "Foreign Keys", "Deskripsi Data"],
    [
        ["users", "id", "-", "Data autentikasi akun (Admin & Mahasiswa), NIM, Prodi, Role"],
        ["kategoris", "id", "-", "Master kategori peralatan lab (Elektronik, Optik, Komputer)"],
        ["barangs", "id", "kategori_id -> kategoris.id", "Master data alat, kode barang unik, stok total, stok tersedia"],
        ["peminjamans", "id", "user_id -> users.id\nbarang_id -> barangs.id", "Tabel transaksi peminjaman, tanggal, kondisi fisik, denda"]
    ]
)

add_h2("3.3 Mekanisme Keamanan & Middleware")
add_p(
    "Untuk memastikan keamanan akses, sistem dilengkapi komponen keamanan mutakhir:\n"
    "1. RoleMiddleware: Memeriksa session login dan atribut role pada pengguna. Mahasiswa dicegah mengakses endpoint /admin/* dan sebaliknya.\n"
    "2. Password Hashing (Bcrypt): Seluruh kata sandi pengguna dienkripsi dengan standar industri bcrypt sebelum disimpan.\n"
    "3. CSRF Protection: Setiap form transaksi dilengkapi token CSRF guna mencegah eksploitasi Cross-Site Request Forgery."
)

# ==========================================
# BAB IV: IMPLEMENTASI STRUKTUR FILE
# ==========================================
add_h1("BAB IV: IMPLEMENTASI STRUKTUR FILE (SESUAI REFERENSI)")

add_p(
    "Susunan file dalam project dibangun secara presisi mengikuti struktur referensi yang diajukan oleh kelompok:"
)

tbl_files = doc.add_table(rows=1, cols=3)
style_table(
    tbl_files,
    [2.2, 1.8, 2.4],
    ["Lokasi Berkas", "Jenis Komponen", "Fungsi dan Tanggung Jawab"],
    [
        ["app/Http/Controllers/Auth/AuthController.php", "Controller", "Menangani proses login multi-role dan logout sesi."],
        ["app/Http/Controllers/Admin/DashboardAdminController.php", "Controller", "Menghitung statistik total aset, denda, dan transaksi aktif."],
        ["app/Http/Controllers/Admin/BarangController.php", "Controller", "Operasi CRUD data barang laboratorium beserta stok."],
        ["app/Http/Controllers/Admin/KategoriController.php", "Controller", "Operasi CRUD kategori peralatan lab."],
        ["app/Http/Controllers/Admin/PeminjamanAdminController.php", "Controller", "Memverifikasi pengembalian, cek kerusakan, hitung denda otomatis."],
        ["app/Http/Controllers/User/DashboardUserController.php", "Controller", "Dashboard mahasiswa, statistik aktif, notifikasi deadline."],
        ["app/Http/Controllers/User/PeminjamanUserController.php", "Controller", "Katalog barang mahasiswa dan form peminjaman mandiri."],
        ["app/Http/Middleware/RoleMiddleware.php", "Middleware", "Guard otoritas akses URL berdasarkan role admin / user."],
        ["app/Models/*.php", "Model (Eloquent)", "Entitas User, Barang, Kategori, dan Peminjaman."],
        ["database/migrations/*.php", "Migrations", "Definisi DDL skema database terversi."],
        ["resources/views/admin/*", "Blade Views", "Antarmuka admin panel (dashboard, barang, kategori, peminjaman)."],
        ["resources/views/user/*", "Blade Views", "Antarmuka portal mahasiswa (dashboard, katalog barang, peminjaman)."],
        ["routes/web.php", "Routing", "Definisi seluruh endpoint URL terproteksi middleware."]
    ]
)

# ==========================================
# BAB V: PENGUJIAN SISTEM & SKENARIO UJI
# ==========================================
add_h1("BAB V: PENGUJIAN SISTEM (TEST CASES & HASIL UJI)")

add_p(
    "Pengujian dilakukan berdasarkan skenario data transaksi nyata (sebagaimana tercantum dalam database seeder):"
)

tbl_uji = doc.add_table(rows=1, cols=6)
style_table(
    tbl_uji,
    [1.4, 0.9, 0.9, 1.0, 1.1, 1.1],
    ["Skenario Uji", "Rencana", "Aktual", "Kondisi", "Hari Telat", "Hasil Denda"],
    [
        ["Kasus 1: Tepat Waktu", "03-10-2026", "03-10-2026", "Baik", "0 hari", "Rp 0 (Lulus)"],
        ["Kasus 2: Telat 5 Hari & Rusak", "05-10-2026", "10-10-2026", "Rusak", "5 hari", "Rp 20.000 (Lulus)"],
        ["Kasus 3: Toleransi 1-3 Hari", "05-10-2026", "07-10-2026", "Baik", "2 hari", "Rp 0 (Toleransi)"],
        ["Kasus 4: Sedang Dipinjam", "20-10-2026", "-", "-", "-", "Stok berkurang (Lulus)"]
    ]
)

add_p(
    "Analisa Hasil Kasus 2 (Terlambat 5 Hari + Kondisi Rusak):\n"
    "• Denda keterlambatan: 5 hari × Rp2.000 = Rp10.000 (karena keterlambatan > 3 hari).\n"
    "• Biaya perbaikan alat: Rp10.000 (karena kondisi = rusak).\n"
    "• Total tagihan denda: Rp10.000 + Rp10.000 = Rp20.000. Sistem menghasilkan output yang persis dan akurat."
)

add_h2("5.2 Kredensial Akun Pengujian (Demo)")
tbl_akun = doc.add_table(rows=1, cols=4)
style_table(
    tbl_akun,
    [1.5, 1.8, 1.4, 1.7],
    ["Peran", "Username / NIM", "Password", "Nama Pemilik Akun"],
    [
        ["Admin", "admin", "admin123", "Petugas Laboratorium"],
        ["Mahasiswa 1", "24051130084", "123456", "Lian Rosyid Alamzah"],
        ["Mahasiswa 2", "24051130094", "123456", "Aditya Danu Marizta"],
        ["Mahasiswa 3", "24051130103", "123456", "Muhammad Ghafur Amanu Putra Warnada"]
    ]
)

# ==========================================
# BAB VI: PEMBAGIAN TUGAS KELOMPOK
# ==========================================
add_h1("BAB VI: PEMBAGIAN TUGAS KELOMPOK")

tbl_tugas = doc.add_table(rows=1, cols=3)
style_table(
    tbl_tugas,
    [1.8, 1.4, 3.2],
    ["Nama Anggota", "NIM", "Rincian Tugas dan Kontribusi Modul"],
    [
        [
            "Lian Rosyid Alamzah",
            "24051130084",
            "• Merancang Skema Database Migrations & Model Eloquent (Kategori, Barang).\n"
            "• Mengembangkan Modul Admin: BarangController dan KategoriController.\n"
            "• Membangun Views Manajemen Barang (Index, Create, Edit) dan Kategori."
        ],
        [
            "Aditya Danu Marizta",
            "24051130094",
            "• Mengimplementasikan aturan bisnis denda keterlambatan dan biaya perbaikan.\n"
            "• Mengembangkan Modul Admin: PeminjamanAdminController.\n"
            "• Membangun Views Transaksi Peminjaman dan Halaman Detail Pengembalian Alat."
        ],
        [
            "Muhammad Ghafur Amanu Putra Warnada",
            "24051130103",
            "• Mengembangkan Modul Autentikasi (AuthController) dan RoleMiddleware.\n"
            "• Membangun Modul User: DashboardUserController dan PeminjamanUserController.\n"
            "• Mendesain Layout Blade, Katalog Barang Mahasiswa, dan Styling Antarmuka."
        ]
    ]
)

# ==========================================
# BAB VII: PANDUAN PENGOPERASIAN
# ==========================================
add_h1("BAB VII: PANDUAN PENGOPERASIAN SISTEM")

add_p(
    "1. Memastikan Layanan XAMPP Aktif:\n"
    "Buka XAMPP Control Panel, lalu jalankan modul Apache dan MySQL hingga berstatus hijau (Running).\n\n"
    "2. Akses Melalui Web Browser:\n"
    "Buka browser (Google Chrome, Microsoft Edge, atau Mozilla Firefox), kemudian ketikkan alamat:\n"
    "👉 http://localhost/UTS%20MSI%20Kelompok%206/\n"
    "Sistem secara otomatis akan mengarahkan pengguna ke halaman Login (/public/login).\n\n"
    "3. Masuk Menggunakan Akun Demo:\n"
    "Gunakan tombol 'Gunakan' pada kotak akun demo untuk login instan sebagai Admin atau Mahasiswa, lalu uji seluruh fitur sistem."
)

# Simpan Dokumen
output_filename = "Laporan_Lengkap_UTS_MSI_Kelompok_6.docx"
doc.save(output_filename)
print(f"File berhasil dibuat: {output_filename}")
